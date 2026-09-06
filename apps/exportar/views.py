"""Exportación a Excel (openpyxl).

`xlsx` recibe un JSON {titulo, hojas:[{nombre, encabezados, filas}]} y devuelve un
archivo .xlsx. Cualquier módulo puede usar este mismo endpoint para exportar sus
resultados.
"""
import io
import json

from django.http import HttpResponse, JsonResponse
from django.shortcuts import render
from django.views.decorators.csrf import csrf_exempt
from django.views.decorators.http import require_POST


def index(request):
    return render(request, 'exportar/index.html')


@csrf_exempt
@require_POST
def xlsx(request):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill

        data = json.loads(request.body or '{}')
        titulo = str(data.get('titulo', 'PlanEstrat'))
        hojas = data.get('hojas', [])
        if not hojas:
            return JsonResponse({'error': 'No hay hojas para exportar.'}, status=400)

        wb = Workbook()
        wb.remove(wb.active)
        for h in hojas:
            nombre = (str(h.get('nombre', 'Hoja'))[:31]) or 'Hoja'
            ws = wb.create_sheet(title=nombre)
            enc = h.get('encabezados', [])
            if enc:
                ws.append(enc)
                for c in ws[1]:
                    c.font = Font(bold=True, color='FFFFFF')
                    c.fill = PatternFill('solid', fgColor='2C5C86')
                    c.alignment = Alignment(vertical='center')
            for fila in h.get('filas', []):
                ws.append(fila)
            for col in ws.columns:
                largo = max((len(str(c.value)) for c in col if c.value is not None), default=10)
                ws.column_dimensions[col[0].column_letter].width = min(max(largo + 2, 10), 55)

        bio = io.BytesIO()
        wb.save(bio)
        bio.seek(0)
        resp = HttpResponse(
            bio.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        safe = ''.join(ch for ch in titulo if ch.isalnum() or ch in ' -_').strip() or 'PlanEstrat'
        resp['Content-Disposition'] = 'attachment; filename="%s.xlsx"' % safe
        return resp
    except Exception:
        return JsonResponse({'error': 'No se pudo generar el Excel.'}, status=400)
